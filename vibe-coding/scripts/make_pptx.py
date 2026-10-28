"""Genera dist/vibe-coding.pptx: 16 diapositivas 16:9 con la imagen 4K de cada
diapositiva a sangre completa y el guion docente en las notas del presentador.

Requiere haber ejecutado antes `npm run build` y `npm run export`.
Uso: python3 scripts/make_pptx.py
"""
import io
import json
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.util import Emu, Inches

ROOT = Path(__file__).resolve().parent.parent
DIST = ROOT / "dist"
SLIDES = ROOT / "build" / "slides"


def fmt_time(seconds: int) -> str:
    return f"{seconds // 60}:{seconds % 60:02d}"


def notes_text(n: dict, by_id: dict) -> str:
    parts = []
    head = "ANEXO · fuera de los 25 minutos" if n.get("annex") else f"DIAPOSITIVA {n['n']} DE 15 · TIEMPO {fmt_time(n['seconds'])}"
    parts.append(f"{head}\n{n['title']}")
    parts.append("GUION ORAL\n" + "\n\n".join(n["guion"]))
    if n.get("tecnica"):
        parts.append("EXPLICACIÓN TÉCNICA\n" + n["tecnica"])
    if n.get("analogia"):
        parts.append("ANALOGÍA\n" + n["analogia"])
    if n.get("pregunta"):
        parts.append("PREGUNTA DE PARTICIPACIÓN\n" + n["pregunta"])
    if n.get("transicion"):
        parts.append("TRANSICIÓN\n" + n["transicion"])
    if n.get("evitar"):
        parts.append("ERRORES CONCEPTUALES QUE EVITAR\n" + "\n".join(f"• {e}" for e in n["evitar"]))
    if n.get("rigor"):
        parts.append("RIGOR: HECHO E INTERPRETACIÓN\n" + n["rigor"])
    if n.get("fuentes"):
        parts.append("FUENTES\n" + "\n".join(f"• {by_id[f]['cite']} {by_id[f]['url']}" for f in n["fuentes"]))
    return "\n\n".join(parts)


def main() -> None:
    data = json.loads((DIST / "notes.json").read_text(encoding="utf-8"))
    notes, sources = data["notes"], data["sources"]
    by_id = {s["id"]: s for s in sources}

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    prs.core_properties.title = "Vibe Coding"
    prs.core_properties.subject = "De una idea escrita a una aplicación funcional"
    prs.core_properties.author = "Top Learning"
    prs.core_properties.keywords = "vibe coding; inteligencia artificial; n8n; Top Learning"

    for n in notes:
        png = SLIDES / f"slide-{n['n']:02d}.png"
        if not png.exists():
            raise SystemExit(f"Falta {png}. Ejecuta primero: npm run export")
        img = Image.open(png).convert("RGB")
        buf = io.BytesIO()
        img.save(buf, "JPEG", quality=90, optimize=True, progressive=True)
        buf.seek(0)
        slide = prs.slides.add_slide(blank)
        pic = slide.shapes.add_picture(buf, 0, 0, width=prs.slide_width, height=prs.slide_height)
        # Texto alternativo para lectores de pantalla.
        pic._element.nvPicPr.cNvPr.set("descr", f"Diapositiva {n['n']}: {n['title']}")
        pic.name = n["title"][:60]
        slide.notes_slide.notes_text_frame.text = notes_text(n, by_id)

    out = DIST / "vibe-coding.pptx"
    prs.save(out)
    print(f"✓ {out.relative_to(ROOT)} ({len(notes)} diapositivas, {out.stat().st_size / 1e6:.1f} MB)")


if __name__ == "__main__":
    main()
